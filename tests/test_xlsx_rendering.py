import datetime
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from engineering_assistant.identity import PLACEHOLDER_DISPLAY_NAME, PLACEHOLDER_STUDENT_ID
from engineering_assistant.xlsx_rendering import MAX_CELL_TEXT, render_solution_xlsx


class _AdvancingClock(datetime.datetime):
    """A UTC clock that moves one second forward on every read.

    openpyxl's writer stamps ``dcterms:modified`` with ``datetime.now`` during
    ``save`` and serializes it at second precision, so two renders differ only
    when their clock reads straddle a second boundary.  Relying on real time
    makes a regression test for that roughly 10% effective; advancing the clock
    deliberately makes it 100%.
    """

    _reads = 0

    @classmethod
    def now(cls, tz=None):
        _AdvancingClock._reads += 1
        base = datetime.datetime(2024, 5, 5, 12, 0, 0, tzinfo=datetime.timezone.utc)
        moment = base + datetime.timedelta(seconds=_AdvancingClock._reads)
        return moment if tz is not None else moment.replace(tzinfo=None)


class XlsxRenderingTests(unittest.TestCase):
    def solution(self):
        return {
            "document_type": "technical_report",
            "title": "Heat Exchanger Results",
            "metadata": {
                "course": "ENGR 328",
                "section": "A",
                "date": "[Date]",
                "student_name": "Private Person",
                "author": "Private Person",
            },
            "abstract": "Measured and calculated heat-transfer results.",
            "sections": [
                {
                    "heading": "Results / Discussion: [Final]",
                    "body": "The measured outlet temperature agreed with the expected trend.",
                    "equations": ["=m_dot*cp*(T_out-T_in)"],
                    "requirement_ids": ["r1", "r2"],
                    "tables": [
                        {
                            "caption": "Table 1: Results?",
                            "headers": ["Quantity", "Value", "Unit", "=Unsafe Header"],
                            "rows": [
                                ["Heat rate", 10.2, "kW", "=1+1"],
                                ["Effectiveness", 0.82, "dimensionless", "+CMD"],
                            ],
                        },
                        {
                            "caption": "Table 1: Results?",
                            "headers": ["Reading", "reading"],
                            "rows": [["Outlet", 42.1]],
                        },
                    ],
                }
            ],
            "calculations": [
                {
                    "id": "energy-balance",
                    "expression": "=m_dot*cp*delta_t",
                    "variables": {"cp": 4.18, "delta_t": 10, "m_dot": 0.25},
                    "expected": 10.45,
                    "tolerance": 0.01,
                    "unit": "kW",
                }
            ],
            "unresolved": [],
        }

    def test_workbook_structure_and_engineering_values_parse_back(self):
        with tempfile.TemporaryDirectory() as directory:
            output = render_solution_xlsx(self.solution(), Path(directory) / "results.xlsx")
            workbook = load_workbook(output, data_only=False)

            self.assertEqual(workbook.sheetnames, [
                "Summary", "Table 1- Results-", "Table 1- Results- (2)", "Calculations"
            ])
            self.assertEqual(workbook["Summary"]["A1"].value, "Heat Exchanger Results")
            self.assertEqual(workbook["Summary"].freeze_panes, "A4")
            results = workbook["Table 1- Results-"]
            self.assertEqual(results.freeze_panes, "A5")
            self.assertEqual(results["A5"].value, "Heat rate")
            self.assertEqual(results["B5"].value, 10.2)
            self.assertEqual(results["C5"].value, "kW")
            self.assertEqual(list(results.tables), ["EngineeringTable1"])
            duplicate_headers = workbook["Table 1- Results- (2)"]
            self.assertEqual(duplicate_headers["A4"].value, "Reading")
            self.assertEqual(duplicate_headers["B4"].value, "reading (2)")

            calculations = workbook["Calculations"]
            self.assertEqual(calculations["A5"].value, "energy-balance")
            self.assertEqual(calculations["C5"].value, '{"cp":4.18,"delta_t":10,"m_dot":0.25}')
            self.assertEqual(calculations["D5"].value, 10.45)
            self.assertEqual(calculations["E5"].value, 0.01)
            self.assertEqual(calculations["F5"].value, "kW")

    def test_formula_like_imported_content_remains_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            output = render_solution_xlsx(self.solution(), Path(directory) / "safe.xlsx")
            workbook = load_workbook(output, data_only=False)
            summary = workbook["Summary"]
            equation = next(
                cell for row in summary.iter_rows() for cell in row
                if cell.value == "=m_dot*cp*(T_out-T_in)"
            )
            table = workbook["Table 1- Results-"]
            calculations = workbook["Calculations"]
            for cell in (equation, table["D4"], table["D5"], table["D6"], calculations["B5"]):
                self.assertEqual(cell.data_type, "s")
            self.assertEqual(table["D5"].value, "=1+1")
            self.assertEqual(calculations["B5"].value, "=m_dot*cp*delta_t")

    def test_identity_is_omitted_or_placeholder_and_metadata_is_private(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            without_identity = render_solution_xlsx(self.solution(), root / "without.xlsx")
            workbook = load_workbook(without_identity)
            values = [
                str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows()
                for cell in row if cell.value is not None
            ]
            self.assertNotIn("Private Person", "\n".join(values))
            self.assertNotIn(PLACEHOLDER_DISPLAY_NAME, values)
            self.assertEqual(workbook.properties.creator, "calvin-autonomy")

            with_identity = render_solution_xlsx(
                self.solution(), root / "with.xlsx", include_identity_placeholders=True
            )
            workbook = load_workbook(with_identity)
            values = [
                str(cell.value) for row in workbook["Summary"].iter_rows()
                for cell in row if cell.value is not None
            ]
            self.assertIn(PLACEHOLDER_DISPLAY_NAME, values)
            self.assertIn(PLACEHOLDER_STUDENT_ID, values)
            self.assertNotIn("Private Person", "\n".join(values))

    def test_output_is_byte_stable_and_cell_text_is_bounded(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            solution = self.solution()
            solution["sections"][0]["tables"][0]["rows"][0][0] = "x" * (MAX_CELL_TEXT + 100)
            first = render_solution_xlsx(solution, root / "first.xlsx")
            second = render_solution_xlsx(solution, root / "second.xlsx")
            self.assertEqual(first.read_bytes(), second.read_bytes())
            workbook = load_workbook(first)
            value = workbook["Table 1- Results-"]["A5"].value
            self.assertEqual(len(value), MAX_CELL_TEXT)
            self.assertTrue(value.endswith(" [truncated]"))

    def test_modified_timestamp_is_frozen_in_core_properties(self):
        """Pins the mechanism: no wall-clock time may reach docProps/core.xml.

        ``created`` was always correct because nothing overwrites it; only
        ``modified`` leaked, because openpyxl reassigns it during ``save``.
        """
        with tempfile.TemporaryDirectory() as directory:
            rendered = render_solution_xlsx(self.solution(), Path(directory) / "book.xlsx")
            with zipfile.ZipFile(rendered) as archive:
                core = archive.read("docProps/core.xml").decode("utf-8")
            self.assertIn("<dcterms:modified", core)
            modified = core.split("<dcterms:modified", 1)[1].split(">", 1)[1].split("<", 1)[0]
            created = core.split("<dcterms:created", 1)[1].split(">", 1)[1].split("<", 1)[0]
            self.assertEqual(modified, "2000-01-01T00:00:00Z")
            self.assertEqual(created, "2000-01-01T00:00:00Z")
            self.assertNotIn(str(datetime.datetime.now().year), core)

    def test_output_is_byte_stable_when_the_clock_advances_between_renders(self):
        """Pins the guarantee, without depending on timing luck.

        Without the fix this fails every run rather than one in ten.
        """
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            solution = self.solution()
            with mock.patch.object(datetime, "datetime", _AdvancingClock):
                first = render_solution_xlsx(solution, root / "first.xlsx")
                second = render_solution_xlsx(solution, root / "second.xlsx")
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_unordered_collection_values_are_rejected(self):
        """A set's iteration order is hash-dependent and would reach the bytes."""
        solution = self.solution()
        solution["sections"][0]["tables"][0]["rows"][0][0] = {"b", "a", "c"}
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                render_solution_xlsx(solution, Path(directory) / "book.xlsx")


if __name__ == "__main__":
    unittest.main()
