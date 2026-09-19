import hashlib
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from PIL import Image

from engineering_assistant.rendering import render_text_docx


class DocxRenderingTests(unittest.TestCase):
    def profile(self):
        return {
            "rendering": {
                "body_font_size": 11,
                "line_spacing": 1.25,
                "margin_inches": 0.75,
                "title_page": True,
                "abstract": True,
                "number_body_pages": True,
                "table_captions_above": True,
                "figure_captions_below": True,
            }
        }

    def solution(self):
        return {
            "document_type": "technical_report",
            "title": "Heat Transfer Lab",
            "metadata": {
                "course": "ENGR 328",
                "section": "Section 1",
                "date": "[Date]",
                "author": "Private Person",
            },
            "abstract": "Measured and calculated results are summarized.",
            "sections": [
                {
                    "heading": "Results",
                    "body": "All reported quantities include units.",
                    "equations": ["q = m cp DeltaT"],
                    "tables": [
                        {
                            "caption": "Table 1: Energy balance results.",
                            "headers": ["Quantity", "Value", "Unit"],
                            "rows": [["Heat rate", "10.2", "kW"]],
                        }
                    ],
                    "figures": [
                        {
                            "path": "figures/result.png",
                            "caption": "Figure 1: Measured response.",
                        }
                    ],
                }
            ],
            "appendices": [
                {"heading": "Appendix A: Calculations", "body": "Calculation record."}
            ],
        }

    def test_report_is_editable_structured_private_and_course_styled(self):
        with tempfile.TemporaryDirectory() as directory:
            run = Path(directory) / "run"
            (run / "deliverables").mkdir(parents=True)
            (run / "figures").mkdir()
            Image.new("RGB", (1200, 600), "white").save(run / "figures" / "result.png")
            output = render_text_docx(
                self.solution(),
                run / "deliverables" / "report.docx",
                style_profile=self.profile(),
            )

            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for expected in (
                "Heat Transfer Lab",
                "ENGR 328",
                "[Student Name]",
                "[Student ID]",
                "Abstract",
                "Results",
                "q = m cp DeltaT",
                "Table 1: Energy balance results.",
                "Figure 1: Measured response.",
                "Appendix A: Calculations",
            ):
                self.assertIn(expected, text)
            self.assertNotIn("Private Person", text)
            self.assertEqual(document.core_properties.author, "[Student Name]")
            self.assertEqual(document.core_properties.last_modified_by, "[Student Name]")
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(0, 0).text, "Quantity")
            self.assertEqual(document.tables[0].cell(1, 1).text, "10.2")
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertEqual(len(document.sections), 2)
            for section in document.sections:
                self.assertAlmostEqual(section.left_margin.inches, 0.75)
                self.assertAlmostEqual(section.right_margin.inches, 0.75)
            self.assertAlmostEqual(document.styles["Normal"].font.size.pt, 11)
            self.assertAlmostEqual(document.styles["Normal"].paragraph_format.line_spacing, 1.25)

            body_xml = document.element.body.xml
            self.assertLess(body_xml.index("Table 1"), body_xml.index("<w:tbl"))
            self.assertLess(body_xml.index("<w:drawing"), body_xml.index("Figure 1"))
            with ZipFile(output) as package:
                xml = b"\n".join(
                    package.read(name) for name in package.namelist() if name.endswith(".xml")
                )
            self.assertIn(b" PAGE ", xml)
            self.assertNotIn(b"Private Person", xml)

    def test_same_input_produces_identical_docx_bytes(self):
        with tempfile.TemporaryDirectory() as directory:
            run = Path(directory) / "run"
            (run / "deliverables").mkdir(parents=True)
            (run / "figures").mkdir()
            Image.new("RGB", (400, 300), "white").save(run / "figures" / "result.png")
            first = render_text_docx(
                self.solution(), run / "deliverables" / "first.docx", style_profile=self.profile()
            )
            second = render_text_docx(
                self.solution(), run / "deliverables" / "second.docx", style_profile=self.profile()
            )
            self.assertEqual(
                hashlib.sha256(first.read_bytes()).hexdigest(),
                hashlib.sha256(second.read_bytes()).hexdigest(),
            )

    def test_profile_and_figure_safety_are_enforced(self):
        with tempfile.TemporaryDirectory() as directory:
            run = Path(directory) / "run"
            (run / "deliverables").mkdir(parents=True)
            missing_abstract = self.solution()
            missing_abstract.pop("abstract")
            with self.assertRaisesRegex(ValueError, "requires an abstract"):
                render_text_docx(
                    missing_abstract,
                    run / "deliverables" / "missing.docx",
                    style_profile=self.profile(),
                )

            unsafe = self.solution()
            unsafe["figures"] = []
            unsafe["sections"][0]["figures"][0]["path"] = "../outside.png"
            with self.assertRaisesRegex(ValueError, "relative to the assignment run"):
                render_text_docx(
                    unsafe,
                    run / "deliverables" / "unsafe.docx",
                    style_profile=self.profile(),
                )


if __name__ == "__main__":
    unittest.main()
