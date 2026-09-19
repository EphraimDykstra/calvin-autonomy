import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from engineering_assistant.artifact_rendering import render_declared_artifacts


class ArtifactRenderingTests(unittest.TestCase):
    def test_all_declared_editable_and_pdf_artifacts_are_rendered(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "run"
            solution = {
                "document_type": "worked_problem",
                "title": "Energy Balance",
                "sections": [
                    {
                        "heading": "Result",
                        "body": "The calculated heat rate is 10 kW.",
                        "requirement_ids": ["r1"],
                        "tables": [
                            {
                                "caption": "Table 1: Result.",
                                "headers": ["Quantity", "Value", "Unit"],
                                "rows": [["Heat rate", 10, "kW"]],
                            }
                        ],
                    }
                ],
                "calculations": [
                    {
                        "id": "r1-check",
                        "expression": "m*cp*dT",
                        "variables": {"m": 1, "cp": 2, "dT": 5},
                        "expected": 10,
                        "tolerance": 1e-9,
                        "unit": "kW",
                    }
                ],
                "unresolved": [],
            }
            plan = {
                "deliverables": [
                    {"path": "deliverables/submission.pdf", "format": "pdf"},
                    {"path": "deliverables/submission.docx", "format": "docx"},
                    {"path": "deliverables/results.xlsx", "format": "xlsx"},
                ]
            }
            outputs = render_declared_artifacts(solution, plan, root)
            self.assertEqual([path.suffix for path in outputs], [".pdf", ".docx", ".xlsx"])
            self.assertEqual(len(PdfReader(str(outputs[0])).pages), 1)
            self.assertIn("Energy Balance", "\n".join(p.text for p in Document(outputs[1]).paragraphs))
            workbook = load_workbook(outputs[2], data_only=False)
            self.assertEqual(workbook["Summary"]["A1"].value, "Energy Balance")
            self.assertEqual(workbook["Calculations"]["D5"].value, 10)

    def test_all_declarations_are_validated_before_any_output_is_written(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "run"
            plan = {
                "deliverables": [
                    {"path": "deliverables/submission.pdf", "format": "pdf"},
                    {"path": "deliverables/results.csv", "format": "csv"},
                ]
            }
            with self.assertRaisesRegex(ValueError, "unsupported render format"):
                render_declared_artifacts({"title": "Result", "sections": []}, plan, root)
            self.assertFalse((root / "deliverables" / "submission.pdf").exists())

            plan["deliverables"][1] = {"path": "../outside.docx", "format": "docx"}
            with self.assertRaises(ValueError):
                render_declared_artifacts({"title": "Result", "sections": []}, plan, root)

            plan["deliverables"][1] = {"path": "deliverables/result.pdf", "format": "docx"}
            with self.assertRaisesRegex(ValueError, "extension does not match"):
                render_declared_artifacts({"title": "Result", "sections": []}, plan, root)


if __name__ == "__main__":
    unittest.main()
