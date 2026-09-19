"""Deterministic, editable XLSX rendering for engineering solutions."""

from __future__ import annotations

import io
import json
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .identity import PLACEHOLDER_DISPLAY_NAME, PLACEHOLDER_STUDENT_ID


MAX_CELL_TEXT = 32_767
_TRUNCATION_MARKER = " [truncated]"
_INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
# openpyxl serializes core properties as naive UTC at second precision
# (packaging/core.py), so the frozen text is derived from _FIXED_TIME rather
# than written out by hand: the two cannot drift apart.
_FIXED_TIME_TEXT = (_FIXED_TIME.replace(tzinfo=None).isoformat(timespec="seconds") + "Z").encode("ascii")
_CORE_PROPERTIES_MEMBER = "docProps/core.xml"
_MODIFIED_ELEMENT = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")

_DARK_BLUE = "1F4E78"
_MEDIUM_BLUE = "5B9BD5"
_LIGHT_BLUE = "D9EAF7"
_LIGHT_GRAY = "E7E6E6"
_WHITE = "FFFFFF"
_THIN_GRAY = Side(style="thin", color="A6A6A6")


def _bounded_text(value: Any) -> str:
    """Return legal Excel text with a visible marker when a cell is clipped."""
    text = _ILLEGAL_XML_CHARACTERS.sub("\ufffd", str(value))
    if len(text) <= MAX_CELL_TEXT:
        return text
    return text[: MAX_CELL_TEXT - len(_TRUNCATION_MARKER)] + _TRUNCATION_MARKER


def _write_literal(cell: Cell, value: Any) -> None:
    """Write imported text as a string, never as an executable formula."""
    text = _bounded_text(value)
    cell.value = text
    cell.data_type = "s"
    if text.lstrip().startswith(("=", "+", "-", "@")):
        cell.quotePrefix = True


def _write_value(cell: Cell, value: Any) -> None:
    """Preserve JSON numeric values while treating every string as literal."""
    if value is None:
        cell.value = None
    elif isinstance(value, bool):
        cell.value = value
    elif isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("workbook values must be finite")
        cell.value = value
    elif isinstance(value, (set, frozenset)):
        # str() on an unordered collection exposes hash-dependent iteration
        # order, which would reach the rendered bytes and break byte stability
        # for any caller that passes one.  Rejecting matches how this function
        # already treats non-finite floats, and how the dict/list branch already
        # rejects nested non-JSON values.
        raise ValueError("workbook values must be ordered; convert sets to a sorted list first")
    elif isinstance(value, (dict, list)):
        _write_literal(cell, json.dumps(value, sort_keys=True, separators=(",", ":")))
    else:
        _write_literal(cell, value)


def _sheet_name(value: Any, used: set[str]) -> str:
    candidate = _INVALID_SHEET_CHARACTERS.sub("-", _bounded_text(value)).strip().strip("'")
    candidate = candidate[:31].rstrip() or "Sheet"
    base = candidate
    suffix = 2
    while candidate.casefold() in used:
        addition = f" ({suffix})"
        candidate = base[: 31 - len(addition)].rstrip() + addition
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _unique_headers(headers: list[Any], column_count: int) -> list[str]:
    """Return nonblank, case-insensitively unique Excel table headers."""
    result: list[str] = []
    used: set[str] = set()
    for index in range(column_count):
        raw = headers[index] if index < len(headers) else ""
        base = _bounded_text(raw).strip() or f"Column {index + 1}"
        candidate = base
        suffix = 2
        while candidate.casefold() in used:
            addition = f" ({suffix})"
            candidate = base[: MAX_CELL_TEXT - len(addition)].rstrip() + addition
            suffix += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def _set_title(cell: Cell, value: Any) -> None:
    _write_literal(cell, value)
    cell.font = Font(name="Aptos Display", size=16, bold=True, color=_WHITE)
    cell.fill = PatternFill("solid", fgColor=_DARK_BLUE)
    cell.alignment = Alignment(vertical="center")


def _set_heading(cell: Cell, value: Any) -> None:
    _write_literal(cell, value)
    cell.font = Font(name="Aptos", size=11, bold=True, color=_WHITE)
    cell.fill = PatternFill("solid", fgColor=_MEDIUM_BLUE)
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _set_label(cell: Cell, value: Any) -> None:
    _write_literal(cell, value)
    cell.font = Font(name="Aptos", size=10, bold=True, color="000000")
    cell.fill = PatternFill("solid", fgColor=_LIGHT_BLUE)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _body_chunks(value: Any) -> list[str]:
    """Keep long prose lossless by spreading it over bounded worksheet rows."""
    text = _ILLEGAL_XML_CHARACTERS.sub("\ufffd", str(value or ""))
    if not text:
        return []
    return [text[index : index + MAX_CELL_TEXT] for index in range(0, len(text), MAX_CELL_TEXT)]


def _finish_sheet(sheet, *, landscape: bool = False, repeat_rows: str | None = None) -> None:
    sheet.freeze_panes = "A4" if sheet.title == "Summary" else "A5"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.35
    sheet.page_margins.right = 0.35
    sheet.page_margins.top = 0.55
    sheet.page_margins.bottom = 0.55
    sheet.oddFooter.center.text = "Page &P of &N"
    if repeat_rows:
        sheet.print_title_rows = repeat_rows


def _set_widths(sheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = max(8, min(float(width), 55))


def _freeze_modified_timestamp(data: bytes) -> bytes:
    """Replace the wall-clock modification time written during ``save``.

    ``workbook.properties.modified`` cannot be frozen before saving: openpyxl's
    writer reassigns it to ``datetime.now`` while serializing, so the real clock
    reaches ``docProps/core.xml``.  Serialization is to second precision, so two
    renders of identical input produced identical bytes only when both clock
    reads landed in the same second -- a byte-stability guarantee that silently
    depended on timing.  This pass runs after ``save`` and is the only place
    that sees the final bytes.
    """
    return _MODIFIED_ELEMENT.sub(rb"\g<1>" + _FIXED_TIME_TEXT + rb"\g<2>", data)


def _canonical_xlsx_bytes(workbook: Workbook) -> bytes:
    """Save stable XML and ZIP metadata so equal inputs yield equal bytes."""
    raw = io.BytesIO()
    workbook.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()), "r")
    canonical = io.BytesIO()
    with source, zipfile.ZipFile(
        canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for original in source.infolist():
            info = zipfile.ZipInfo(original.filename, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = original.external_attr
            content = source.read(original.filename)
            if original.filename == _CORE_PROPERTIES_MEMBER:
                content = _freeze_modified_timestamp(content)
            target.writestr(info, content, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    return canonical.getvalue()


def render_solution_xlsx(
    solution: dict[str, Any],
    output: Path,
    *,
    include_identity_placeholders: bool = False,
) -> Path:
    """Render the existing solution schema as an editable engineering workbook.

    Identity fields are omitted by default. When a template requires them, only
    neutral placeholders are written; reusable workbooks never inherit a name
    or identifier from arbitrary solution metadata.
    """
    if not isinstance(solution, dict):
        raise ValueError("solution must be an object")
    sections = solution.get("sections", [])
    calculations = solution.get("calculations", [])
    if not isinstance(sections, list):
        raise ValueError("solution sections must be a list")
    if not isinstance(calculations, list):
        raise ValueError("solution calculations must be a list")

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.properties.creator = "calvin-autonomy"
    workbook.properties.lastModifiedBy = "calvin-autonomy"
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    workbook.properties.title = _bounded_text(solution.get("title", "Engineering Assignment"))
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False

    used_names = {"summary"}
    summary = workbook.active
    summary.title = "Summary"
    summary.merge_cells("A1:D1")
    _set_title(summary["A1"], solution.get("title", "Engineering Assignment"))
    summary.row_dimensions[1].height = 27
    summary_row = 3

    metadata = solution.get("metadata", {})
    metadata = metadata if isinstance(metadata, dict) else {}
    approved_metadata = (
        ("Document Type", solution.get("document_type", "engineering_workbook")),
        ("Course", metadata.get("course")),
        ("Section", metadata.get("section")),
        ("Date", metadata.get("date")),
    )
    for label, value in approved_metadata:
        if value in (None, ""):
            continue
        _set_label(summary.cell(summary_row, 1), label)
        _write_literal(summary.cell(summary_row, 2), value)
        summary.merge_cells(start_row=summary_row, start_column=2, end_row=summary_row, end_column=4)
        summary_row += 1
    if include_identity_placeholders:
        for label, value in (
            ("Student", PLACEHOLDER_DISPLAY_NAME),
            ("Student ID", PLACEHOLDER_STUDENT_ID),
        ):
            _set_label(summary.cell(summary_row, 1), label)
            _write_literal(summary.cell(summary_row, 2), value)
            summary.merge_cells(start_row=summary_row, start_column=2, end_row=summary_row, end_column=4)
            summary_row += 1

    abstract = solution.get("abstract")
    if abstract:
        summary_row += 1
        _set_heading(summary.cell(summary_row, 1), "Abstract")
        summary.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
        summary_row += 1
        for chunk in _body_chunks(abstract):
            _write_literal(summary.cell(summary_row, 1), chunk)
            summary.cell(summary_row, 1).alignment = Alignment(vertical="top", wrap_text=True)
            summary.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
            summary_row += 1

    table_number = 0
    for section_number, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            raise ValueError("every solution section must be an object")
        heading = section.get("heading") or f"Section {section_number}"
        summary_row += 1
        _set_heading(summary.cell(summary_row, 1), heading)
        summary.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
        summary_row += 1
        requirement_ids = section.get("requirement_ids", [])
        if isinstance(requirement_ids, list) and requirement_ids:
            _set_label(summary.cell(summary_row, 1), "Requirements")
            _write_literal(summary.cell(summary_row, 2), ", ".join(str(item) for item in requirement_ids))
            summary.merge_cells(start_row=summary_row, start_column=2, end_row=summary_row, end_column=4)
            summary_row += 1
        for chunk in _body_chunks(section.get("body", "")):
            _write_literal(summary.cell(summary_row, 1), chunk)
            summary.cell(summary_row, 1).alignment = Alignment(vertical="top", wrap_text=True)
            summary.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
            summary_row += 1
        equations = section.get("equations", [])
        if equations is not None and not isinstance(equations, list):
            raise ValueError("section equations must be a list")
        for equation in equations or []:
            _set_label(summary.cell(summary_row, 1), "Equation")
            _write_literal(summary.cell(summary_row, 2), equation)
            summary.merge_cells(start_row=summary_row, start_column=2, end_row=summary_row, end_column=4)
            summary_row += 1

        tables = section.get("tables", [])
        if tables is not None and not isinstance(tables, list):
            raise ValueError("section tables must be a list")
        for table in tables or []:
            if not isinstance(table, dict):
                raise ValueError("every section table must be an object")
            headers = table.get("headers", [])
            rows = table.get("rows", [])
            if not isinstance(headers, list) or not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
                raise ValueError("table headers and rows must be lists")
            table_number += 1
            caption = table.get("caption") or f"{heading} Table {table_number}"
            table_sheet = workbook.create_sheet(_sheet_name(caption, used_names))
            column_count = max([len(headers), *(len(row) for row in rows)], default=1)
            normalized_headers = _unique_headers(headers, column_count)
            table_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
            _set_title(table_sheet.cell(1, 1), caption)
            _set_label(table_sheet.cell(2, 1), "Section")
            _write_literal(table_sheet.cell(2, 2), heading)
            for column, header in enumerate(normalized_headers, start=1):
                _set_heading(table_sheet.cell(4, column), header)
            for row_number, row in enumerate(rows, start=5):
                for column in range(1, column_count + 1):
                    cell = table_sheet.cell(row_number, column)
                    _write_value(cell, row[column - 1] if column <= len(row) else None)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = Border(bottom=_THIN_GRAY)
            if rows:
                excel_table = Table(
                    displayName=f"EngineeringTable{table_number}",
                    ref=f"A4:{table_sheet.cell(4 + len(rows), column_count).coordinate}",
                )
                excel_table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                table_sheet.add_table(excel_table)
            _set_widths(table_sheet, {
                get_column_letter(column): max(
                    len(_bounded_text(normalized_headers[column - 1])) + 2,
                    *(min(len(_bounded_text(row[column - 1])), 50) + 2 for row in rows if column <= len(row)),
                    12,
                )
                for column in range(1, column_count + 1)
            })
            _finish_sheet(table_sheet, landscape=column_count > 4, repeat_rows="1:4")

    _set_widths(summary, {"A": 22, "B": 25, "C": 25, "D": 25})
    _finish_sheet(summary, repeat_rows="1:1")

    if calculations:
        calculation_sheet = workbook.create_sheet(_sheet_name("Calculations", used_names))
        columns = ("ID", "Expression", "Variables", "Expected", "Tolerance", "Unit")
        calculation_sheet.merge_cells("A1:F1")
        _set_title(calculation_sheet["A1"], "Calculation Records")
        for column, header in enumerate(columns, start=1):
            _set_heading(calculation_sheet.cell(4, column), header)
        for row_number, record in enumerate(calculations, start=5):
            if not isinstance(record, dict):
                raise ValueError("every calculation record must be an object")
            values = (
                record.get("id", ""),
                record.get("expression", ""),
                record.get("variables", {}),
                record.get("expected"),
                record.get("tolerance", 1e-8),
                record.get("unit", ""),
            )
            for column, value in enumerate(values, start=1):
                cell = calculation_sheet.cell(row_number, column)
                _write_value(cell, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=_THIN_GRAY)
        calculation_table = Table(
            displayName="CalculationRecords",
            ref=f"A4:F{4 + len(calculations)}",
        )
        calculation_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        calculation_sheet.add_table(calculation_table)
        _set_widths(calculation_sheet, {
            "A": 20, "B": 34, "C": 42, "D": 16, "E": 16, "F": 18,
        })
        _finish_sheet(calculation_sheet, landscape=True, repeat_rows="1:4")

    workbook.active = 0
    output.write_bytes(_canonical_xlsx_bytes(workbook))
    return output
